from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Tiêu đề chính ──────────────────────────────────────────────────────────
title = doc.add_heading('BÀI KIỂM TRA GIỮA KỲ', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_heading('KIẾN TRÚC VÀ THIẾT KẾ PHẦN MỀM', 1)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Môn học: Kiến Trúc Và Thiết Kế Phần Mềm').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Họ và tên: Nguyễn Văn A').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('MSSV: XXXXXXXX').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# CÂU 1 — MỞ RỘNG THIẾT KẾ
# ══════════════════════════════════════════════════════════════════════════════
doc.add_heading('CÂU 1 — MỞ RỘNG THIẾT KẾ', 1)

# ── 1.1 Sơ đồ ERD ──────────────────────────────────────────────────────────
doc.add_heading('1.1 Sơ Đồ ERD — Module Quản Lý Khuyến Mãi', 2)

doc.add_paragraph(
    'Bổ sung module Quản Lý Khuyến Mãi vào hệ thống bán trái cây. '
    'Sơ đồ ERD dưới đây mô tả các bảng dữ liệu mới và mối quan hệ với bảng Đơn Hàng (DON_HANG).'
)

erd = doc.add_paragraph()
erd.style = 'No Spacing'
run = erd.add_run(
"""
Bảng KHUYEN_MAI (mới):
┌─────────────────────────────────────────────────────────┐
│                     KHUYEN_MAI                           │
├─────────────────────────────────────────────────────────┤
│ PK  MaKhuyenMai     INT                                 │
│     MaVoucher       NVARCHAR(50)   UNIQUE               │
│     TenKhuyenMai    NVARCHAR(200)                       │
│     LoaiKhuyenMai   NVARCHAR(30)                        │
│        → GiamTheoTienMat / GiamTheoPhantram / MuaNTangM │
│     GiaTriGiam      DECIMAL(18,2)                       │
│     DonHangToiThieu DECIMAL(18,2)                       │
│     NgayBatDau      DATETIME                            │
│     NgayKetThuc     DATETIME                            │
└─────────────────────────────────────────────────────────┘

Bảng DON_HANG (bổ sung khóa ngoại):
┌───────────────────────────────┐
│         DON_HANG              │
├───────────────────────────────┤
│ PK  MaDonHang     INT         │
│ FK  MaKhachHang   INT         │
│ FK  MaVoucher     NVARCHAR  ◄─┼── liên kết KHUYEN_MAI (nullable)
│     NgayDatHang   DATETIME    │
│     TrangThai     NVARCHAR    │
│     TongTienGoc   DECIMAL     │
│     SoTienGiam    DECIMAL     │
│     TongTienThanhToan DECIMAL │
└───────────────────────────────┘

Quan hệ:
  KHUYEN_MAI  1 ──── N  DON_HANG
  (Một mã voucher có thể dùng cho nhiều đơn hàng, nullable)

  KHACH_HANG  1 ──── N  DON_HANG
  DON_HANG    1 ──── N  CHI_TIET_DON_HANG
  SAN_PHAM    1 ──── N  CHI_TIET_DON_HANG
"""
)
run.font.name = 'Courier New'
run.font.size = Pt(9)

# ── 1.2 Interface tính giá ──────────────────────────────────────────────────
doc.add_heading('1.2 Interface Dịch Vụ Tính Giá', 2)

doc.add_paragraph(
    'Interface ITinhGiaService được thiết kế để có thể áp dụng nhiều loại mã giảm giá khác nhau '
    'mà không cần sửa code hiện có (tuân thủ Open/Closed Principle).'
)

iface = doc.add_paragraph()
run2 = iface.add_run(
"""
«interface»
ITinhGiaService
──────────────────────────────────────────────────────
+ TinhSoTienGiam(tongTienGoc: decimal,
                 khuyenMai: KhuyenMai) : decimal
  → Trả về số tiền được giảm

+ TinhTongTienSauGiam(tongTienGoc: decimal,
                      khuyenMai: KhuyenMai?) : decimal
  → Trả về tổng tiền sau khi áp dụng khuyến mãi
          ▲
          │ implements
          │
   TinhGiaService
──────────────────────────────────────────────────────
  Hỗ trợ 3 loại chiết khấu:
  • GiamTheoTienMat  → giảm số tiền cố định (VNĐ)
  • GiamTheoPhantram → giảm theo % (có áp trần 100.000 VNĐ)
  • MuaNTangM        → tặng hàng tương đương số tiền nhất định
"""
)
run2.font.name = 'Courier New'
run2.font.size = Pt(9)

doc.add_paragraph(
    'Lý do dùng Interface: Khi cần thêm loại chiết khấu mới (ví dụ: giảm theo thành viên VIP), '
    'chỉ cần tạo class mới implement ITinhGiaService mà không ảnh hưởng code hiện tại.'
)

# ══════════════════════════════════════════════════════════════════════════════
# CÂU 2 — TRIỂN KHAI XÁC NHẬN ĐẶT HÀNG
# ══════════════════════════════════════════════════════════════════════════════
doc.add_heading('CÂU 2 — TRIỂN KHAI CHỨC NĂNG XÁC NHẬN ĐẶT HÀNG', 1)

doc.add_heading('2.1 Design Pattern — Builder Pattern', 2)
doc.add_paragraph(
    'Áp dụng Builder Pattern để tạo đối tượng DonHang. '
    'Đối tượng đơn hàng có nhiều thuộc tính tùy chọn (voucher, nhiều sản phẩm), '
    'Builder Pattern giúp tạo đơn hàng từng bước theo kiểu method chaining, '
    'code dễ đọc và tránh constructor quá nhiều tham số.'
)

builder_code = doc.add_paragraph()
r = builder_code.add_run(
"""// Cách sử dụng Builder Pattern
var donHang = new DonHangBuilder()
    .VoiKhachHang(101, "Nguyễn Văn An", "123 Lê Lợi, TP.HCM")
    .ThemSanPham(maSanPham: 1, soLuong: 2)   // 2 kg Xoài
    .ThemSanPham(maSanPham: 3, soLuong: 3)   // 3 kg Bưởi
    .ApDungVoucher("GIAM50K")
    .Build();
"""
)
r.font.name = 'Courier New'
r.font.size = Pt(9)

doc.add_heading('2.2 Kiến Trúc Phân Tầng', 2)
arch = doc.add_paragraph()
r2 = arch.add_run(
"""
┌──────────────────────────────────────────────────┐
│              TẦNG TRÌNH BÀY                       │
│   Program.cs — Console, nhận input, in kết quả   │
└────────────────────┬─────────────────────────────┘
                     │ gọi (DI)
┌────────────────────▼─────────────────────────────┐
│              TẦNG NGHIỆP VỤ                       │
│   DonHangService  — luồng xác nhận đặt hàng      │
│   TinhGiaService  — tính giá sau khuyến mãi      │
└──────┬──────────────────────┬────────────────────┘
       │ DI                   │ DI
┌──────▼──────┐  ┌────────────▼────┐  ┌────────────┐
│ TẦNG DỮ LIỆU│  │ TẦNG DỮ LIỆU   │  │TẦNG DỮ LI. │
│SanPhamRepo. │  │ DonHangRepo.    │  │KhuyenMaiR. │
└─────────────┘  └─────────────────┘  └────────────┘
"""
)
r2.font.name = 'Courier New'
r2.font.size = Pt(9)

doc.add_heading('2.3 Luồng Xử Lý Xác Nhận Đặt Hàng', 2)
flow = doc.add_paragraph()
r3 = flow.add_run(
"""
Người Dùng → DonHangBuilder.Build() → DonHang
                                           │
                        DonHangService.XacNhanDatHang()
                                           │
          ┌────────────┬──────────────────┬┴───────────────┐
          ▼            ▼                  ▼                 ▼
    Validate      Đồng bộ giá       Kiểm tra          Áp dụng
    đầu vào      từ Repository      Tồn Kho           Voucher
          │            │                  │                 │
          └────────────┴──────────────────┴─────────────────┘
                                           │ (tất cả OK)
                                    Lưu Đơn Hàng
                                           │
                                    Cập Nhật Tồn Kho
                                           │
                                    Trả Kết Quả
"""
)
r3.font.name = 'Courier New'
r3.font.size = Pt(9)

# ══════════════════════════════════════════════════════════════════════════════
# CÂU 3 — TỐI ƯU HÓA — EXCEPTION HANDLING
# ══════════════════════════════════════════════════════════════════════════════
doc.add_heading('CÂU 3 — TỐI ƯU HÓA KIẾN TRÚC — XỬ LÝ NGOẠI LỆ', 1)

doc.add_heading('3.1 Phân Cấp Ngoại Lệ Tùy Chỉnh', 2)
exc = doc.add_paragraph()
r4 = exc.add_run(
"""
Exception  (hệ thống .NET)
    └── NgoaiLeUngDung              ← lỗi gốc của ứng dụng
            ├── DonHangRongException
            │      → đơn hàng không có sản phẩm nào
            ├── SanPhamKhongTonTaiException
            │      → mã sản phẩm không tồn tại
            ├── TonKhoKhongDuException
            │      → số lượng yêu cầu > tồn kho
            └── VoucherKhongHopLeException
                   → voucher không tồn tại / hết hạn / chưa đủ điều kiện
"""
)
r4.font.name = 'Courier New'
r4.font.size = Pt(9)

doc.add_heading('3.2 Chiến Lược Xử Lý Trong Service', 2)
exc2 = doc.add_paragraph()
r5 = exc2.add_run(
"""
public DonHang XacNhanDatHang(DonHang donHang)
{
    try
    {
        KiemTraDonHangHopLe(donHang);   // throw DonHangRongException
        DongBoGiaSanPham(donHang);      // throw SanPhamKhongTonTaiException
        KiemTraTonKho(donHang);         // throw TonKhoKhongDuException
        ApDungKhuyenMai(donHang);       // throw VoucherKhongHopLeException
        var donHangDaLuu = _donHangRepo.Luu(donHang);
        CapNhatTonKho(donHang);
        return donHangDaLuu;
    }
    catch (NgoaiLeUngDung ex)
    {
        // Lỗi nghiệp vụ đã biết → log + ném lại cho tầng trên
        Console.WriteLine($"[LỖI NGHIỆP VỤ] {ex.Message}");
        throw;
    }
    catch (Exception ex)
    {
        // Lỗi hệ thống không mong đợi → bọc lại, che thông tin nội bộ
        throw new NgoaiLeUngDung("Lỗi hệ thống khi xử lý đặt hàng.", ex);
    }
}
"""
)
r5.font.name = 'Courier New'
r5.font.size = Pt(9)

doc.add_heading('3.3 Nguyên Tắc Áp Dụng', 2)

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = 'Nguyên Tắc'
hdr[1].text = 'Áp Dụng Trong Dự Án'

rows_data = [
    ('Fail Fast', 'Validate đơn hàng rỗng ngay đầu hàm'),
    ('Specific Exception', 'Mỗi lỗi có class riêng với thông tin chi tiết'),
    ('No Silent Catch', 'Không bao giờ catch rồi bỏ qua — luôn log hoặc throw'),
    ('Wrap External Errors', 'Lỗi ngoài bọc bởi NgoaiLeUngDung'),
    ('User-Friendly Message', 'Thông điệp lỗi tiếng Việt rõ ràng cho người dùng'),
]
for principle, apply in rows_data:
    row = table.add_row().cells
    row[0].text = principle
    row[1].text = apply

doc.save('/home/user/Hoi_thao/MASV_NGUYENVANA/ThietKe.docx')
print("Tạo file docx thành công!")
