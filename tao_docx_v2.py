from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Căn lề trang ─────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

def h1(text):
    p = doc.add_heading(text, 1)
    p.runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)
    return p

def h2(text):
    p = doc.add_heading(text, 2)
    p.runs[0].font.color.rgb = RGBColor(0x2e, 0x86, 0xab)
    return p

def body(text):
    p = doc.add_paragraph(text)
    p.runs[0].font.size = Pt(11)
    return p

def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(text)
    r.font.name = 'Courier New'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    # shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EBF5FB')
    pPr.append(shd)
    return p

def add_img(path, width_inches=6.0, caption_text=''):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width_inches))
    if caption_text:
        cap = doc.add_paragraph(caption_text)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.runs[0]
        r.font.size  = Pt(9)
        r.font.italic = True
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph()

def sep():
    doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════
# TRANG BÌA
# ══════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('BÀI KIỂM TRA GIỮA KỲ\nKIẾN TRÚC VÀ THIẾT KẾ PHẦN MỀM')
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)

sep()
for line in ['Họ và tên:   Nguyễn Văn A',
             'MSSV:        XXXXXXXX',
             'Lớp:         KTTKPM - XX',
             'Thời gian:   48 giờ']:
    lp = doc.add_paragraph(line)
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.runs[0].font.size = Pt(12)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════
# CÂU 1 — MỞ RỘNG THIẾT KẾ
# ══════════════════════════════════════════════════════════════════════════
h1('CÂU 1 — MỞ RỘNG THIẾT KẾ')

# ── 1.1 ERD ──────────────────────────────────────────────────────────────
h2('1.1 Sơ Đồ ERD — Module Quản Lý Khuyến Mãi')
body(
    'Để mở rộng hệ thống bán trái cây với module Quản Lý Khuyến Mãi, '
    'ta bổ sung bảng KHUYEN_MAI và thêm khóa ngoại MaVoucher vào bảng DON_HANG. '
    'Sơ đồ ERD bên dưới thể hiện đầy đủ các bảng dữ liệu và mối quan hệ:'
)
sep()
add_img('/home/user/Hoi_thao/diagrams/erd.png', width_inches=6.5,
        caption_text='Hình 1.1 — Sơ đồ ERD hệ thống bán trái cây (bảng xanh lá = bảng mới)')

body('Mô tả các mối quan hệ chính:')
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for cell, txt in zip(table.rows[0].cells, ['Bảng A', 'Quan Hệ', 'Bảng B']):
    cell.text = txt
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)

for a, rel, b in [
    ('KHACH_HANG', '1 ──── N', 'DON_HANG'),
    ('DON_HANG',   '1 ──── N', 'CHI_TIET_DON_HANG'),
    ('SAN_PHAM',   '1 ──── N', 'CHI_TIET_DON_HANG'),
    ('KHUYEN_MAI', '1 ──── N (nullable)', 'DON_HANG'),
]:
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = a, rel, b

sep()

# ── 1.2 Interface ─────────────────────────────────────────────────────────
h2('1.2 Interface Dịch Vụ Tính Giá — ITinhGiaService')
body(
    'Interface ITinhGiaService được thiết kế theo nguyên tắc Open/Closed: '
    'thêm loại chiết khấu mới không cần sửa code hiện có, '
    'chỉ cần mở rộng TinhGiaService hoặc tạo class mới implement interface này.'
)
sep()
add_img('/home/user/Hoi_thao/diagrams/interface_tinh_gia.png', width_inches=6.2,
        caption_text='Hình 1.2 — Sơ đồ class ITinhGiaService và các thành phần liên quan')

body('Hỗ trợ 3 loại chiết khấu được định nghĩa trong enum LoaiKhuyenMai:')
for item in [
    '• GiamTheoTienMat  — Giảm số tiền cố định (VNĐ), không vượt tổng tiền đơn hàng.',
    '• GiamTheoPhantram — Giảm theo phần trăm, có áp trần tối đa 100.000 VNĐ.',
    '• MuaNTangM        — Tặng hàng hóa tương đương giá trị nhất định.',
]:
    p = doc.add_paragraph(item)
    p.runs[0].font.size = Pt(11)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════
# CÂU 2 — XÁC NHẬN ĐẶT HÀNG
# ══════════════════════════════════════════════════════════════════════════
h1('CÂU 2 — TRIỂN KHAI CHỨC NĂNG XÁC NHẬN ĐẶT HÀNG (C#)')

# ── 2.1 Builder Pattern ───────────────────────────────────────────────────
h2('2.1 Design Pattern — Builder Pattern')
body(
    'Áp dụng Builder Pattern để tạo đối tượng DonHang. '
    'Đối tượng đơn hàng có nhiều thuộc tính tùy chọn (danh sách sản phẩm, voucher…). '
    'Builder Pattern cho phép xây dựng đơn hàng từng bước theo kiểu Method Chaining, '
    'code dễ đọc như ngôn ngữ tự nhiên và tránh constructor quá nhiều tham số.'
)
code_block(
"""// ── Sử dụng DonHangBuilder (Builder Pattern) ──────────────────────
var donHang = new DonHangBuilder()
    .VoiKhachHang(101, "Nguyễn Văn An", "123 Lê Lợi, TP.HCM")
    .ThemSanPham(maSanPham: 1, soLuong: 2)   // 2 kg Xoài Cát Hòa Lộc
    .ThemSanPham(maSanPham: 3, soLuong: 3)   // 3 kg Bưởi Da Xanh
    .ApDungVoucher("GIAM50K")
    .Build();

// ── Gọi service xác nhận đặt hàng ─────────────────────────────────
var ketQua = donHangSvc.XacNhanDatHang(donHang);
ketQua.HienThiDonHang();"""
)

# ── 2.2 Kiến trúc phân tầng ───────────────────────────────────────────────
h2('2.2 Kiến Trúc Phân Tầng')
body(
    'Hệ thống được chia thành 4 tầng rõ ràng theo mô hình N-Layer, '
    'kết hợp Repository Pattern và Dependency Injection để tách biệt các mối quan tâm.'
)
sep()
add_img('/home/user/Hoi_thao/diagrams/kien_truc_phan_tang.png', width_inches=6.2,
        caption_text='Hình 2.1 — Kiến trúc phân tầng với Repository Pattern và DI')

# ── 2.3 Luồng xử lý ──────────────────────────────────────────────────────
h2('2.3 Luồng Xử Lý Xác Nhận Đặt Hàng')
body('Luồng xử lý trong DonHangService.XacNhanDatHang() trải qua 6 bước tuần tự:')
for step in [
    '1. Validate đầu vào — kiểm tra đơn hàng không rỗng, số lượng > 0.',
    '2. Đồng bộ giá từ Repository — lấy giá chính thức từ CSDL, không tin giá từ client.',
    '3. Kiểm tra tồn kho — nếu không đủ → ném TonKhoKhongDuException.',
    '4. Áp dụng khuyến mãi — validate voucher và tính số tiền giảm.',
    '5. Lưu đơn hàng — ghi vào Repository với trạng thái DaXacNhan.',
    '6. Cập nhật tồn kho — trừ số lượng sản phẩm đã đặt.',
]:
    p = doc.add_paragraph(step)
    p.runs[0].font.size = Pt(11)

sep()
add_img('/home/user/Hoi_thao/diagrams/luong_dat_hang.png', width_inches=6.5,
        caption_text='Hình 2.2 — Luồng xử lý xác nhận đặt hàng (Sequence Diagram)')

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════
# CÂU 3 — TỐI ƯU HÓA — EXCEPTION HANDLING
# ══════════════════════════════════════════════════════════════════════════
h1('CÂU 3 — TỐI ƯU HÓA KIẾN TRÚC — XỬ LÝ NGOẠI LỆ')

h2('3.1 Phân Cấp Ngoại Lệ Tùy Chỉnh')
body(
    'Thay vì dùng Exception chung, hệ thống định nghĩa phân cấp ngoại lệ riêng. '
    'Mỗi lỗi nghiệp vụ có class riêng chứa đầy đủ thông tin ngữ cảnh.'
)
code_block(
"""Exception  (.NET system)
    └── NgoaiLeUngDung              ← lỗi gốc của ứng dụng
            ├── DonHangRongException
            │      → đơn hàng không có sản phẩm nào
            ├── SanPhamKhongTonTaiException
            │      → mã sản phẩm không tồn tại (maSanPham)
            ├── TonKhoKhongDuException
            │      → yêu cầu {soLuong} > tồn kho {hienCo}
            └── VoucherKhongHopLeException
                   → voucher không tồn tại / hết hạn / chưa đủ điều kiện"""
)

h2('3.2 Chiến Lược Xử Lý Trong Service Layer')
code_block(
"""public DonHang XacNhanDatHang(DonHang donHang)
{
    try
    {
        KiemTraDonHangHopLe(donHang);   // DonHangRongException
        DongBoGiaSanPham(donHang);      // SanPhamKhongTonTaiException
        KiemTraTonKho(donHang);         // TonKhoKhongDuException
        ApDungKhuyenMai(donHang);       // VoucherKhongHopLeException
        var donHangDaLuu = _donHangRepo.Luu(donHang);
        CapNhatTonKho(donHang);
        return donHangDaLuu;
    }
    catch (NgoaiLeUngDung ex)
    {
        // Lỗi nghiệp vụ đã biết → log thông báo và ném lại cho tầng trên
        Console.WriteLine($"[LỖI NGHIỆP VỤ] {ex.Message}");
        throw;
    }
    catch (Exception ex)
    {
        // Lỗi hệ thống không mong đợi → bọc lại để che thông tin nội bộ
        throw new NgoaiLeUngDung("Lỗi hệ thống khi xử lý đặt hàng.", ex);
    }
}"""
)

h2('3.3 Nguyên Tắc Áp Dụng')
table2 = doc.add_table(rows=1, cols=2)
table2.style = 'Table Grid'
for cell, txt in zip(table2.rows[0].cells, ['Nguyên Tắc', 'Áp Dụng Trong Dự Án']):
    cell.text = txt
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)

for principle, apply_txt in [
    ('Fail Fast',             'Validate đơn hàng rỗng ngay đầu hàm, không chờ đến bước lưu'),
    ('Specific Exception',    'Mỗi lỗi có class riêng chứa thông tin ngữ cảnh đầy đủ'),
    ('No Silent Catch',       'Không bao giờ catch rồi bỏ qua — luôn log hoặc throw'),
    ('Wrap External Errors',  'Lỗi từ bên ngoài được bọc bởi NgoaiLeUngDung'),
    ('User-Friendly Message', 'Thông điệp lỗi tiếng Việt, rõ ràng cho người dùng cuối'),
]:
    row = table2.add_row().cells
    row[0].text, row[1].text = principle, apply_txt

sep()

# ── Quyết định kiến trúc ──────────────────────────────────────────────────
h1('QUYẾT ĐỊNH KIẾN TRÚC')
for title, content in [
    ('Tại Sao Dùng Repository Pattern?',
     'Repository Pattern tạo lớp trừu tượng giữa tầng nghiệp vụ và dữ liệu. '
     'Khi cần chuyển từ dữ liệu in-memory sang SQL Server thực, '
     'chỉ cần tạo class repository mới implement cùng interface — '
     'không thay đổi bất kỳ dòng code nào trong Service.'),
    ('Tại Sao Dùng Dependency Injection?',
     'DonHangService nhận các dependency qua constructor thay vì tự khởi tạo. '
     'Điều này giúp dễ viết unit test (mock dependency), '
     'tuân thủ Dependency Inversion Principle, '
     'và dễ thay thế implementation (ví dụ: đổi sang Redis cache).'),
    ('Tại Sao Dùng Builder Thay Vì Constructor?',
     'DonHang có nhiều thuộc tính tùy chọn. '
     'Nếu dùng constructor sẽ cần nhiều overload hoặc tham số quá nhiều (khó đọc, dễ nhầm thứ tự). '
     'Builder cho phép tạo đơn hàng từng bước, dễ đọc và mở rộng.'),
]:
    h2(title)
    body(content)
    sep()

doc.save('/home/user/Hoi_thao/MASV_NGUYENVANA/ThietKe.docx')
print('✔ Tạo Word file thành công!')
